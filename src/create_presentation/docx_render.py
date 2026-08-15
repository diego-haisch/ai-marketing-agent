"""Render the CV as an editable Word document (.docx) with ApplyChain tokens.

Word cannot reproduce gradients or exact card radius, so the design is
simplified to solid token colors: dark teal header band, honey section
underlines, teal table headers, and light highlights.
"""

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

from . import cv_data
from .tokens import COLORS as C

FONT = "Montserrat"


# ---------------------------------------------------------------------------
# Low-level helpers
# ---------------------------------------------------------------------------

def _shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def _border(p, edge, color, sz=16, space=4):
    pPr = p._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), str(sz))
    el.set(qn("w:space"), str(space))
    el.set(qn("w:color"), color)
    pBdr.append(el)


def _shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _cell_margins(cell, top=80, start=110, bottom=80, end=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = OxmlElement(f"w:{tag}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        tcMar.append(el)
    tcPr.append(tcMar)


def _run(p, text, size=9, bold=False, color=None, caps=False, italic=False):
    r = p.add_run(text.upper() if caps else text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    rPr = r._element.get_or_add_rPr()
    rFonts = rPr.rFonts
    if rFonts is not None:
        rFonts.set(qn("w:cs"), FONT)
        rFonts.set(qn("w:eastAsia"), FONT)
    if color:
        r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return r


def _para(container, after=2, before=0):
    p = container.add_paragraph() if hasattr(container, "add_paragraph") else container
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    return p


# ---------------------------------------------------------------------------
# Sections
# ---------------------------------------------------------------------------

def _header(doc):
    p = _para(doc, after=1)
    _shade_paragraph(p, C["egyptian_teal"])
    _run(p, "MICRO1 · AI DATA LAB", size=8, bold=True, color=C["bayou"], caps=True)

    p = _para(doc, after=1)
    _shade_paragraph(p, C["egyptian_teal"])
    _run(p, cv_data.PROFILE["name"], size=24, bold=True, color=C["white"])

    p = _para(doc, after=2)
    _shade_paragraph(p, C["egyptian_teal"])
    _run(p, cv_data.PROFILE["role"], size=10.5, bold=True, color=C["honey"], caps=True)

    p = _para(doc, after=3)
    _shade_paragraph(p, C["egyptian_teal"])
    _run(p, cv_data.PROFILE["tagline"], size=9.5, color="#C9E4E8")

    p = _para(doc, after=4)
    _shade_paragraph(p, C["egyptian_teal"])
    _run(p, " ", size=5)
    _border(p, "bottom", C["honey"], sz=32, space=2)

    contact = list(cv_data.PROFILE["contact"].items())
    table = doc.add_table(rows=2, cols=2)
    table.autofit = False
    for i, (label, value) in enumerate(contact):
        cell = table.cell(i // 2, i % 2)
        cell.width = Cm(9.0)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
        _shade_cell(cell, C["egyptian_teal"])
        _cell_margins(cell)
        cell.text = ""
        pl = _para(cell.paragraphs[0], after=0)
        _run(pl, label, size=6.5, bold=True, color=C["bayou"], caps=True)
        pv = cell.add_paragraph()
        pv.paragraph_format.space_after = Pt(0)
        _run(pv, value, size=8.5, color=C["white"])

    _para(doc, after=6)


def _section_title(doc, text, border_color=C["honey"]):
    p = _para(doc, after=5, before=2)
    _run(p, text, size=11.5, bold=True, color=C["teal_blue"], caps=True)
    _border(p, "bottom", border_color, sz=16, space=3)
    return p


def _summary(doc):
    _section_title(doc, "Professional Summary")
    p = _para(doc, after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    _run(p, cv_data.PROFILE["summary"], size=9)


def _role(doc, role):
    p = _para(doc, after=1, before=2)
    pf = p.paragraph_format
    pf.tab_stops.add_tab_stop(Cm(18.0), WD_TAB_ALIGNMENT.RIGHT)
    _run(p, role["role"], size=10, bold=True, color=C["dark_space"])
    _run(p, f'  ·  {role["company"]}', size=10, bold=True, color=C["egyptian_teal"])
    _run(p, f'\t{role["dates"]}', size=8.5, color=C["honey_dark"])

    for b in role["bullets"]:
        p = _para(doc, after=1)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.55)
        pf.first_line_indent = Cm(-0.55)
        _run(p, "\u25B8  ", size=8.4, bold=True, color=C["bayou"])
        _run(p, b, size=8.4)


def _experience(doc):
    _section_title(doc, "Professional Experience")
    for idx, role in enumerate(cv_data.EXPERIENCE):
        if idx:
            _para(doc, after=4)
        _role(doc, role)
    _para(doc, after=6)


def _highlight(doc):
    _section_title(doc, cv_data.AI_SECTION["title"], border_color=C["egyptian_teal"])
    for b in cv_data.AI_SECTION["bullets"]:
        p = _para(doc, after=2)
        pf = p.paragraph_format
        pf.left_indent = Cm(0.55)
        pf.first_line_indent = Cm(-0.55)
        _shade_paragraph(p, C["chart_area"])
        _border(p, "left", C["egyptian_teal"], sz=24, space=6)
        _run(p, "\u25B8  ", size=8.6, bold=True, color=C["egyptian_teal"])
        _run(p, b, size=8.6)
    _para(doc, after=6)


def _two_col_table(doc, left_cells, right_cells, widths=(9.0, 9.0)):
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    left.width, right.width = Cm(widths[0]), Cm(widths[1])
    _cell_margins(left, end=220)
    _cell_margins(right, start=220)
    left.text, right.text = "", ""
    for cell, content in ((left, left_cells), (right, right_cells)):
        content(cell)


def _competency_group(cell, group):
    p = _para(cell.paragraphs[0], after=1)
    _run(p, group["name"], size=7.5, bold=True, color=C["teal_blue"], caps=True)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    _run(p, "  ·  ".join(group["items"]), size=7.5, color=C["dark_space"])


def _competencies(doc):
    _section_title(doc, "Core Competencies")
    groups = cv_data.COMPETENCY_GROUPS

    def left(cell):
        for g in groups[:4]:
            _competency_group(cell, g)

    def right(cell):
        for g in groups[4:]:
            _competency_group(cell, g)

    _two_col_table(doc, left, right)
    _para(doc, after=6)


def _education(doc):
    p = _para(doc, after=2, before=2)
    _run(p, "EDUCATION", size=8.5, bold=True, color=C["egyptian_teal"], caps=True)
    for edu in cv_data.EDUCATION:
        p = _para(doc, after=1)
        _run(p, edu["degree"], size=8.5, bold=True, color=C["dark_space"])
        p = _para(doc, after=0)
        _run(p, edu["school"], size=8, color=C["egyptian_teal"])
        _run(p, f'  ·  {edu["years"]}', size=7.5, color="#7A8A90")


def _languages(doc):
    p = _para(doc, after=2, before=2)
    _run(p, "LANGUAGES", size=8.5, bold=True, color=C["egyptian_teal"], caps=True)
    for lang in cv_data.LANGUAGES:
        p = _para(doc, after=1)
        _run(p, lang["name"], size=8.5, bold=True, color=C["dark_space"])
        _run(p, f'  ·  {lang["level"]}  ({lang["pct"]}%)', size=8, color=C["egyptian_teal"])


def _skills(doc):
    p = _para(doc, after=2, before=2)
    _run(p, "IT SKILLS", size=8.5, bold=True, color=C["egyptian_teal"], caps=True)
    p = _para(doc, after=4)
    _run(p, "  ·  ".join(cv_data.SKILLS), size=8)


def _background(doc):
    _section_title(doc, "Education, Languages & IT Skills")

    def left(cell):
        _education(cell)
        _languages(cell)

    def right(cell):
        _skills(cell)

    _two_col_table(doc, left, right)
    _para(doc, after=6)


def _fit_table(doc, section):
    p = _para(doc, after=2, before=2)
    _run(p, section["title"], size=10, bold=True, color=C["teal_blue"])

    rows = 1 + len(section["rows"])
    table = doc.add_table(rows=rows, cols=2)
    table.style = "Table Grid"
    table.autofit = False

    headers = ("micro1 requirement", "Evidence")
    for j, text in enumerate(headers):
        cell = table.cell(0, j)
        cell.width = Cm(6.4) if j == 0 else Cm(11.6)
        _shade_cell(cell, C["egyptian_teal"])
        _cell_margins(cell, top=40, bottom=40, start=90, end=90)
        cell.paragraphs[0].text = ""
        _run(cell.paragraphs[0], text, size=8, bold=True, color=C["white"])

    for i, (req, evid) in enumerate(section["rows"], start=1):
        for j, value in enumerate((req, evid)):
            cell = table.cell(i, j)
            cell.width = Cm(6.4) if j == 0 else Cm(11.6)
            _cell_margins(cell, top=40, bottom=40, start=90, end=90)
            cell.paragraphs[0].text = ""
            _run(
                cell.paragraphs[0],
                value,
                size=8,
                bold=(j == 0),
                color=C["teal_blue"] if j == 0 else C["dark_space"],
            )
            if i % 2 == 0:
                _shade_cell(cell, C["desert_storm"])
    _para(doc, after=6)


def _fit(doc):
    _section_title(doc, "Position Fit — micro1")
    for section in cv_data.FIT_SECTIONS:
        _fit_table(doc, section)


def _qa_annex(doc):
    doc.add_page_break()
    _section_title(doc, "Application Q&A — micro1 Consulting Domain Expert")
    for qa in cv_data.APPLICATION_QA:
        p = _para(doc, after=2, before=2)
        _run(p, qa["question"], size=9.5, bold=True, color=C["teal_blue"])
        for paragraph in qa["answer"]:
            p = _para(doc, after=2)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            _run(p, paragraph, size=8.5)


# ---------------------------------------------------------------------------
# Document assembly
# ---------------------------------------------------------------------------

def save_docx(output_path: str) -> None:
    doc = Document()

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(1.4)
    section.bottom_margin = Cm(1.4)
    section.left_margin = Cm(1.5)
    section.right_margin = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(9)
    normal.paragraph_format.space_after = Pt(2)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.element.rPr.rFonts.set(qn("w:cs"), FONT)

    doc.core_properties.title = f"CV — {cv_data.PROFILE['name']}"

    _header(doc)
    _summary(doc)
    _experience(doc)
    _highlight(doc)
    _competencies(doc)
    _background(doc)
    _fit(doc)
    _qa_annex(doc)

    doc.save(output_path)
